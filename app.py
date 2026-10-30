
import os
from io import BytesIO

from flask import (
    Flask, render_template, request, redirect,
    url_for, flash, send_file
)

from flask_sqlalchemy import SQLAlchemy

from flask_login import (
    LoginManager, UserMixin, login_user,
    login_required, logout_user, current_user
)

from flask_wtf.csrf import CSRFProtect

from werkzeug.security import (
    generate_password_hash, check_password_hash
)

from dotenv import load_dotenv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

load_dotenv()

app = Flask(__name__)

app.config["SECRET_KEY"] = os.getenv("SECRET_KEY")
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)
csrf = CSRFProtect(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"


# --------------------------------------------------
# DATABASE MODELS
# --------------------------------------------------

# Teacher account table
class Teacher(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    full_name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password = db.Column(db.String(255), nullable=False)


# Class table
class ClassRoom(db.Model):
    id = db.Column(db.Integer, primary_key=True)

    teacher_id = db.Column(
        db.Integer,
        db.ForeignKey("teacher.id"),
        nullable=False
    )

    academic_year = db.Column(db.String(20), nullable=False)
    year_level = db.Column(db.String(30), nullable=False)
    section = db.Column(db.String(50), nullable=False)
    subject = db.Column(db.String(50), nullable=False, default="")


# Student table
class Student(db.Model):
    id = db.Column(db.Integer, primary_key=True)

    class_id = db.Column(
        db.Integer,
        db.ForeignKey("class_room.id"),
        nullable=False
    )

    full_name = db.Column(db.String(150), nullable=False)


# Attendance table
class Attendance(db.Model):
    id = db.Column(db.Integer, primary_key=True)

    student_id = db.Column(
        db.Integer,
        db.ForeignKey("student.id"),
        nullable=False
    )

    attendance_date = db.Column(db.Date, nullable=False)

    status = db.Column(db.String(10), nullable=False)

    __table_args__ = (
        db.UniqueConstraint(
            "student_id",
            "attendance_date",
            name="unique_student_attendance_date"
        ),
    )


# --------------------------------------------------
# LOGIN MANAGEMENT
# --------------------------------------------------

@login_manager.user_loader
def load_teacher(teacher_id):
    return db.session.get(Teacher, int(teacher_id))


# --------------------------------------------------
# HOME PAGE
# --------------------------------------------------

@app.route("/")
def home():
    return render_template("home.html")


# --------------------------------------------------
# TEACHER REGISTRATION
# --------------------------------------------------

@app.route("/register", methods=["GET", "POST"])
def register():

    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == "POST":

        full_name = request.form.get("full_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        if not full_name or not email or not password:
            flash("Please fill in all fields.", "error")
            return render_template("register.html")

        if len(password) < 8:
            flash("Password must be at least 8 characters.", "error")
            return render_template("register.html")

        existing_teacher = db.session.execute(
            db.select(Teacher).filter_by(email=email)
        ).scalar_one_or_none()

        if existing_teacher:
            flash("That email is already registered.", "error")
            return render_template("register.html")

        new_teacher = Teacher(
            full_name=full_name,
            email=email,
            password=generate_password_hash(password)
        )

        db.session.add(new_teacher)
        db.session.commit()

        flash("Account created! You can now log in.", "success")

        return redirect(url_for("login"))

    return render_template("register.html")


# --------------------------------------------------
# TEACHER LOGIN
# --------------------------------------------------

@app.route("/login", methods=["GET", "POST"])
def login():

    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == "POST":

        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        teacher = db.session.execute(
            db.select(Teacher).filter_by(email=email)
        ).scalar_one_or_none()

        if teacher and check_password_hash(teacher.password, password):

            login_user(teacher)

            return redirect(url_for("dashboard"))

        flash("Incorrect email or password.", "error")

    return render_template("login.html")


# --------------------------------------------------
# TEACHER DASHBOARD
# --------------------------------------------------

@app.route("/dashboard")
@login_required
def dashboard():
    return render_template("dashboard.html", teacher=current_user)


# --------------------------------------------------
# VIEW ALL CLASSES
# --------------------------------------------------

@app.route("/classes")
@login_required
def classes():

    teacher_classes = db.session.execute(
        db.select(ClassRoom).filter_by(
            teacher_id=current_user.id
        )
    ).scalars().all()

    return render_template(
        "classes.html",
        classes=teacher_classes
    )


# --------------------------------------------------
# CREATE A CLASS
# --------------------------------------------------

@app.route("/classes/create", methods=["GET", "POST"])
@login_required
def create_class():

    if request.method == "POST":

        academic_year = request.form.get("academic_year", "").strip()
        year_level = request.form.get("year_level", "").strip()
        section = request.form.get("section", "").strip()
        subject = request.form.get("subject", "").strip()

        if not academic_year or not year_level or not section or not subject:
            flash("Please fill in all class fields.", "error")
            return render_template("create_class.html")

        # Validate academic year
        try:
            first_year, second_year = academic_year.split("-")

            if (
                len(first_year) != 4
                or len(second_year) != 4
                or not first_year.isdigit()
                or not second_year.isdigit()
                or int(second_year) != int(first_year) + 1
            ):
                raise ValueError

        except ValueError:

            flash(
                "Please enter a valid academic year, such as 2026-2027.",
                "error"
            )

            return render_template("create_class.html")

        new_class = ClassRoom(
            teacher_id=current_user.id,
            academic_year=academic_year,
            year_level=year_level,
            section=section,
            subject=subject
        )

        db.session.add(new_class)
        db.session.commit()

        flash("Class created successfully!", "success")

        return redirect(url_for("classes"))

    return render_template("create_class.html")


# --------------------------------------------------
# EDIT AN EXISTING CLASS
# --------------------------------------------------

@app.route("/classes/<int:class_id>/edit", methods=["GET", "POST"])
@login_required
def edit_class(class_id):

    class_room = ClassRoom.query.filter_by(
        id=class_id,
        teacher_id=current_user.id
    ).first_or_404()

    if request.method == "POST":

        academic_year = request.form.get("academic_year", "").strip()
        year_level = request.form.get("year_level", "").strip()
        section = request.form.get("section", "").strip()
        subject = request.form.get("subject", "").strip()

        if not academic_year or not year_level or not section or not subject:

            flash("Please fill in all class fields.", "error")

            return render_template(
                "edit_class.html",
                class_room=class_room
            )

        # Validate academic year
        try:
            first_year, second_year = academic_year.split("-")

            if (
                len(first_year) != 4
                or len(second_year) != 4
                or not first_year.isdigit()
                or not second_year.isdigit()
                or int(second_year) != int(first_year) + 1
            ):
                raise ValueError

        except ValueError:

            flash(
                "Please enter a valid academic year, such as 2026-2027.",
                "error"
            )

            return render_template(
                "edit_class.html",
                class_room=class_room
            )

        # Update class information
        class_room.academic_year = academic_year
        class_room.year_level = year_level
        class_room.section = section
        class_room.subject = subject

        db.session.commit()

        flash("Class updated successfully!", "success")

        return redirect(url_for("classes"))

    return render_template(
        "edit_class.html",
        class_room=class_room
    )


# --------------------------------------------------
# DELETE AN EXISTING CLASS
# --------------------------------------------------

@app.route("/classes/<int:class_id>/delete", methods=["POST"])
@login_required
def delete_class(class_id):

    class_room = ClassRoom.query.filter_by(
        id=class_id,
        teacher_id=current_user.id
    ).first_or_404()

    # Get all students belonging to this class
    students = Student.query.filter_by(
        class_id=class_id
    ).all()

    # Get student IDs
    student_ids = [student.id for student in students]

    # Delete attendance records first
    if student_ids:

        Attendance.query.filter(
            Attendance.student_id.in_(student_ids)
        ).delete(synchronize_session=False)

    # Delete students
    Student.query.filter_by(
        class_id=class_id
    ).delete(synchronize_session=False)

    # Delete the class
    db.session.delete(class_room)

    db.session.commit()

    flash(
        "Class and its associated records deleted successfully!",
        "success"
    )

    return redirect(url_for("classes"))


# --------------------------------------------------
# VIEW CLASS AND STUDENTS
# --------------------------------------------------

@app.route("/classes/<int:class_id>")
@login_required
def class_detail(class_id):

    class_room = db.session.execute(
        db.select(ClassRoom).filter_by(
            id=class_id,
            teacher_id=current_user.id
        )
    ).scalar_one_or_none()

    if class_room is None:

        flash("Class not found.", "error")

        return redirect(url_for("classes"))

    students = db.session.execute(
        db.select(Student).filter_by(
            class_id=class_room.id
        )
    ).scalars().all()

    # Sort by surname, then full name
    students.sort(
        key=lambda student: (
            student.full_name.split()[-1].casefold(),
            student.full_name.casefold()
        )
    )

    return render_template(
        "class_detail.html",
        class_room=class_room,
        students=students
    )


# --------------------------------------------------
# ADD ONE STUDENT
# --------------------------------------------------

@app.route("/classes/<int:class_id>/add-student", methods=["POST"])
@login_required
def add_student(class_id):

    class_room = db.session.execute(
        db.select(ClassRoom).filter_by(
            id=class_id,
            teacher_id=current_user.id
        )
    ).scalar_one_or_none()

    if class_room is None:

        flash("Class not found.", "error")

        return redirect(url_for("classes"))

    full_name = request.form.get("full_name", "").strip()

    if not full_name:

        flash("Please enter the student's name.", "error")

        return redirect(
            url_for("class_detail", class_id=class_id)
        )

    existing_students = db.session.execute(
        db.select(Student).filter_by(class_id=class_id)
    ).scalars().all()

    for student in existing_students:

        if student.full_name.casefold() == full_name.casefold():

            flash(
                "That student is already in this class.",
                "error"
            )

            return redirect(
                url_for("class_detail", class_id=class_id)
            )

    new_student = Student(
        class_id=class_id,
        full_name=full_name
    )

    db.session.add(new_student)
    db.session.commit()

    flash("Student added successfully!", "success")

    return redirect(
        url_for("class_detail", class_id=class_id)
    )


# --------------------------------------------------
# ADD MULTIPLE STUDENTS
# --------------------------------------------------

@app.route("/classes/<int:class_id>/bulk-add", methods=["POST"])
@login_required
def bulk_add_students(class_id):

    class_room = db.session.execute(
        db.select(ClassRoom).filter_by(
            id=class_id,
            teacher_id=current_user.id
        )
    ).scalar_one_or_none()

    if class_room is None:

        flash("Class not found.", "error")

        return redirect(url_for("classes"))

    names_text = request.form.get("student_names", "")
    names = names_text.splitlines()

    existing_students = db.session.execute(
        db.select(Student).filter_by(class_id=class_id)
    ).scalars().all()

    existing_names = {
        student.full_name.casefold()
        for student in existing_students
    }

    added_count = 0
    skipped_count = 0

    for name in names:

        full_name = name.strip()

        if not full_name:
            continue

        if full_name.casefold() in existing_names:

            skipped_count += 1
            continue

        new_student = Student(
            class_id=class_id,
            full_name=full_name
        )

        db.session.add(new_student)

        existing_names.add(full_name.casefold())

        added_count += 1

    db.session.commit()

    flash(
        f"Added {added_count} student(s). "
        f"Skipped {skipped_count} duplicate name(s).",
        "success"
    )

    return redirect(
        url_for("class_detail", class_id=class_id)
    )


# --------------------------------------------------
# DOWNLOAD CLASS STUDENT LIST AS EXCEL
# --------------------------------------------------

@app.route("/classes/<int:class_id>/students/export")
@login_required
def export_class_students(class_id):

    # Make sure the class belongs to the logged-in teacher
    class_room = ClassRoom.query.filter_by(
        id=class_id,
        teacher_id=current_user.id
    ).first_or_404()

    # Get students belonging to this class
    students = Student.query.filter_by(
        class_id=class_id
    ).all()

    # Sort students alphabetically by surname
    students.sort(
        key=lambda student: (
            student.full_name.split()[-1].casefold(),
            student.full_name.casefold()
        )
    )

    # Create Excel workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Student List"

    # Class information
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "STUDENT ATTENDANCE MONITORING SYSTEM"
    sheet["A1"].font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )
    sheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="172B4D"
    )
    sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    sheet["A2"] = "Class"
    sheet["B2"] = (
        f"{class_room.year_level} - {class_room.section}"
    )

    sheet["A3"] = "Subject"
    sheet["B3"] = class_room.subject

    sheet["A4"] = "Academic Year"
    sheet["B4"] = class_room.academic_year

    sheet["A5"] = "Total Students"
    sheet["B5"] = len(students)

    # Style class information labels
    for row in range(2, 6):

        sheet.cell(row=row, column=1).font = Font(
            bold=True,
            color="172B4D"
        )

    # Student table headings
    sheet["A7"] = "No."
    sheet["B7"] = "Student Name"

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="172B4D"
    )

    for cell in sheet[7]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Add student records
    for index, student in enumerate(students, start=1):

        row_number = index + 7

        sheet.cell(
            row=row_number,
            column=1,
            value=index
        )

        sheet.cell(
            row=row_number,
            column=2,
            value=student.full_name
        )

        sheet.cell(
            row=row_number,
            column=1
        ).alignment = Alignment(
            horizontal="center"
        )

        # Alternate row colors
        if index % 2 == 0:

            for column in range(1, 3):

                sheet.cell(
                    row=row_number,
                    column=column
                ).fill = PatternFill(
                    fill_type="solid",
                    fgColor="F4F7FB"
                )

    # Add borders to the student table
    thin_border = Border(
        bottom=Side(
            style="thin",
            color="DCE4ED"
        )
    )

    for row in sheet.iter_rows(
        min_row=7,
        max_row=7 + len(students),
        min_col=1,
        max_col=2
    ):

        for cell in row:
            cell.border = thin_border

    # Adjust column widths
    sheet.column_dimensions["A"].width = 15
    sheet.column_dimensions["B"].width = 40

    # Adjust row heights
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[7].height = 22

    # Freeze the heading row
    sheet.freeze_panes = "A8"

    # Save Excel workbook in memory
    excel_file = BytesIO()

    workbook.save(excel_file)

    excel_file.seek(0)

    # Prepare the download filename
    filename = (
        f"{class_room.year_level}_"
        f"{class_room.section}_Students.xlsx"
    )

    filename = filename.replace(" ", "_")

    # Send Excel file to the browser
    return send_file(
        excel_file,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# --------------------------------------------------
# DELETE A STUDENT
# --------------------------------------------------

@app.route(
    "/classes/<int:class_id>/delete-student/<int:student_id>",
    methods=["POST"]
)
@login_required
def delete_student(class_id, student_id):

    class_room = db.session.execute(
        db.select(ClassRoom).filter_by(
            id=class_id,
            teacher_id=current_user.id
        )
    ).scalar_one_or_none()

    if class_room is None:

        flash("Class not found.", "error")

        return redirect(url_for("classes"))

    student = db.session.execute(
        db.select(Student).filter_by(
            id=student_id,
            class_id=class_id
        )
    ).scalar_one_or_none()

    if student is None:

        flash("Student not found.", "error")

    else:

        # Delete the student's attendance records first
        Attendance.query.filter_by(
            student_id=student.id
        ).delete(synchronize_session=False)

        # Delete the student
        db.session.delete(student)

        db.session.commit()

        flash("Student removed successfully.", "success")

    return redirect(
        url_for("class_detail", class_id=class_id)
    )


# --------------------------------------------------
# VIEW AND RECORD ATTENDANCE
# --------------------------------------------------

@app.route(
    "/classes/<int:class_id>/attendance",
    methods=["GET", "POST"]
)
@login_required
def attendance(class_id):

    from datetime import date

    # Check class ownership
    class_room = db.session.execute(
        db.select(ClassRoom).filter_by(
            id=class_id,
            teacher_id=current_user.id
        )
    ).scalar_one_or_none()

    if class_room is None:

        flash("Class not found.", "error")

        return redirect(url_for("classes"))

    # Get selected attendance date
    if request.method == "POST":

        selected_date_text = request.form.get(
            "attendance_date", ""
        )

    else:

        selected_date_text = request.args.get(
            "date",
            date.today().isoformat()
        )

    try:

        selected_date = date.fromisoformat(
            selected_date_text
        )

    except ValueError:

        flash("Please select a valid attendance date.", "error")

        return redirect(
            url_for("attendance", class_id=class_id)
        )

    # Get students in this class
    students = db.session.execute(
        db.select(Student).filter_by(
            class_id=class_id
        )
    ).scalars().all()

    # Sort students alphabetically
    students.sort(
        key=lambda student: (
            student.full_name.split()[-1].casefold(),
            student.full_name.casefold()
        )
    )

    if request.method == "POST":

        allowed_statuses = [
            "Present",
            "Absent",
            "Late"
        ]

        attendance_records = {}

        # Validate every student's attendance
        for student in students:

            status = request.form.get(
                f"status_{student.id}",
                ""
            )

            if status not in allowed_statuses:

                flash(
                    "Please select a valid status for every student.",
                    "error"
                )

                return render_template(
                    "attendance.html",
                    class_room=class_room,
                    students=students,
                    selected_date=selected_date,
                    attendance_records={}
                )

            attendance_records[student.id] = status

        try:

            # Save or update attendance
            for student in students:

                record = db.session.execute(
                    db.select(Attendance).filter_by(
                        student_id=student.id,
                        attendance_date=selected_date
                    )
                ).scalar_one_or_none()

                if record is None:

                    record = Attendance(
                        student_id=student.id,
                        attendance_date=selected_date,
                        status=attendance_records[student.id]
                    )

                    db.session.add(record)

                else:

                    record.status = attendance_records[student.id]

            db.session.commit()

            flash("Attendance saved successfully!", "success")

            return redirect(
                url_for(
                    "attendance",
                    class_id=class_id,
                    date=selected_date.isoformat()
                )
            )

        except Exception:

            db.session.rollback()

            app.logger.exception(
                "Failed to save attendance"
            )

            flash(
                "Unable to save attendance. Please try again.",
                "error"
            )

    # Load saved attendance for selected date
    attendance_records = {}

    for student in students:

        record = db.session.execute(
            db.select(Attendance).filter_by(
                student_id=student.id,
                attendance_date=selected_date
            )
        ).scalar_one_or_none()

        if record:
            attendance_records[student.id] = record.status

    return render_template(
        "attendance.html",
        class_room=class_room,
        students=students,
        selected_date=selected_date,
        attendance_records=attendance_records
    )


# --------------------------------------------------
# ATTENDANCE REPORTS
# --------------------------------------------------

@app.route("/classes/<int:class_id>/attendance-reports")
@login_required
def attendance_reports(class_id):

    from datetime import date

    class_room = db.session.execute(
        db.select(ClassRoom).filter_by(
            id=class_id,
            teacher_id=current_user.id
        )
    ).scalar_one_or_none()

    if class_room is None:

        flash("Class not found.", "error")

        return redirect(url_for("classes"))

    selected_date_text = request.args.get(
        "date",
        date.today().isoformat()
    )

    try:

        selected_date = date.fromisoformat(
            selected_date_text
        )

    except ValueError:

        flash("Please select a valid attendance date.", "error")

        return redirect(
            url_for(
                "attendance_reports",
                class_id=class_id
            )
        )

    students = db.session.execute(
        db.select(Student).filter_by(
            class_id=class_id
        )
    ).scalars().all()

    students.sort(
        key=lambda student: (
            student.full_name.split()[-1].casefold(),
            student.full_name.casefold()
        )
    )

    attendance_records = {}

    for student in students:

        record = db.session.execute(
            db.select(Attendance).filter_by(
                student_id=student.id,
                attendance_date=selected_date
            )
        ).scalar_one_or_none()

        if record:
            attendance_records[student.id] = record.status

    # Count attendance statuses
    present_count = sum(
        1 for status in attendance_records.values()
        if status == "Present"
    )

    absent_count = sum(
        1 for status in attendance_records.values()
        if status == "Absent"
    )

    late_count = sum(
        1 for status in attendance_records.values()
        if status == "Late"
    )

    recorded_count = len(attendance_records)
    total_students = len(students)
    not_recorded_count = total_students - recorded_count

    return render_template(
        "attendance_reports.html",
        class_room=class_room,
        students=students,
        selected_date=selected_date,
        attendance_records=attendance_records,
        present_count=present_count,
        absent_count=absent_count,
        late_count=late_count,
        recorded_count=recorded_count,
        total_students=total_students,
        not_recorded_count=not_recorded_count
    )


# --------------------------------------------------
# DOWNLOAD ATTENDANCE REPORT AS EXCEL
# --------------------------------------------------

@app.route("/classes/<int:class_id>/attendance-reports/export")
@login_required
def export_attendance_report(class_id):

    from datetime import date

    # Check if the class belongs to the logged-in teacher
    class_room = ClassRoom.query.filter_by(
        id=class_id,
        teacher_id=current_user.id
    ).first_or_404()

    # Get the selected report date
    selected_date_text = request.args.get(
        "date",
        date.today().isoformat()
    )

    try:
        selected_date = date.fromisoformat(
            selected_date_text
        )

    except ValueError:
        flash("Please select a valid attendance date.", "error")

        return redirect(
            url_for(
                "attendance_reports",
                class_id=class_id
            )
        )

    # Get students belonging to this class
    students = Student.query.filter_by(
        class_id=class_id
    ).all()

    # Sort students alphabetically by surname
    students.sort(
        key=lambda student: (
            student.full_name.split()[-1].casefold(),
            student.full_name.casefold()
        )
    )

    # Get attendance records for the selected date
    attendance_records = {}

    for student in students:

        record = Attendance.query.filter_by(
            student_id=student.id,
            attendance_date=selected_date
        ).first()

        if record:
            attendance_records[student.id] = record.status

    # Count attendance statuses
    present_count = sum(
        1 for status in attendance_records.values()
        if status == "Present"
    )

    absent_count = sum(
        1 for status in attendance_records.values()
        if status == "Absent"
    )

    late_count = sum(
        1 for status in attendance_records.values()
        if status == "Late"
    )

    recorded_count = len(attendance_records)
    total_students = len(students)
    not_recorded_count = total_students - recorded_count

    # Create Excel workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance Report"

    # Define styles
    navy_fill = PatternFill(
        fill_type="solid",
        fgColor="172B4D"
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="21845A"
    )

    light_fill = PatternFill(
        fill_type="solid",
        fgColor="F4F7FB"
    )

    white_font = Font(
        bold=True,
        color="FFFFFF"
    )

    title_font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )

    label_font = Font(
        bold=True,
        color="172B4D"
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="DCE4ED"
        )
    )

    center_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    # --------------------------------------------------
    # REPORT TITLE
    # --------------------------------------------------

    sheet.merge_cells("A1:C1")
    sheet["A1"] = "STUDENT ATTENDANCE MONITORING SYSTEM"
    sheet["A1"].font = title_font
    sheet["A1"].fill = navy_fill
    sheet["A1"].alignment = center_alignment

    sheet.row_dimensions[1].height = 30

    # --------------------------------------------------
    # CLASS INFORMATION
    # --------------------------------------------------

    sheet["A3"] = "Class"
    sheet["B3"] = (
        f"{class_room.year_level} - {class_room.section}"
    )

    sheet["A4"] = "Subject"
    sheet["B4"] = class_room.subject

    sheet["A5"] = "Academic Year"
    sheet["B5"] = class_room.academic_year

    sheet["A6"] = "Report Date"
    sheet["B6"] = selected_date.strftime("%B %d, %Y")

    for row in range(3, 7):
        sheet.cell(row=row, column=1).font = label_font

    # --------------------------------------------------
    # ATTENDANCE SUMMARY
    # --------------------------------------------------

    sheet["A8"] = "ATTENDANCE SUMMARY"
    sheet.merge_cells("A8:B8")

    sheet["A8"].font = white_font
    sheet["A8"].fill = green_fill
    sheet["A8"].alignment = center_alignment

    summary_data = [
        ("Present", present_count),
        ("Absent", absent_count),
        ("Late", late_count),
        ("Not Recorded", not_recorded_count),
        ("Total Students", total_students)
    ]

    for index, (label, count) in enumerate(
        summary_data,
        start=9
    ):
        sheet.cell(
            row=index,
            column=1,
            value=label
        )

        sheet.cell(
            row=index,
            column=2,
            value=count
        )

        sheet.cell(
            row=index,
            column=1
        ).font = label_font

        sheet.cell(
            row=index,
            column=2
        ).alignment = center_alignment

    # --------------------------------------------------
    # STUDENT ATTENDANCE TABLE
    # --------------------------------------------------

    table_start = 16

    sheet.cell(
        row=table_start,
        column=1,
        value="No."
    )

    sheet.cell(
        row=table_start,
        column=2,
        value="Student Name"
    )

    sheet.cell(
        row=table_start,
        column=3,
        value="Attendance Status"
    )

    for cell in sheet[table_start]:

        cell.font = white_font
        cell.fill = navy_fill
        cell.alignment = center_alignment

    # Add student attendance records
    for index, student in enumerate(students, start=1):

        row_number = table_start + index

        status = attendance_records.get(
            student.id,
            "Not Recorded"
        )

        sheet.cell(
            row=row_number,
            column=1,
            value=index
        )

        sheet.cell(
            row=row_number,
            column=2,
            value=student.full_name
        )

        sheet.cell(
            row=row_number,
            column=3,
            value=status
        )

        sheet.cell(
            row=row_number,
            column=1
        ).alignment = center_alignment

        sheet.cell(
            row=row_number,
            column=3
        ).alignment = center_alignment

        # Alternate row colors
        if index % 2 == 0:

            for column in range(1, 4):

                sheet.cell(
                    row=row_number,
                    column=column
                ).fill = light_fill

        # Add bottom borders
        for column in range(1, 4):

            sheet.cell(
                row=row_number,
                column=column
            ).border = thin_border

    # Add borders to table headings
    for column in range(1, 4):

        sheet.cell(
            row=table_start,
            column=column
        ).border = thin_border

    # --------------------------------------------------
    # EXCEL FORMATTING
    # --------------------------------------------------

    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 35
    sheet.column_dimensions["C"].width = 25

    sheet.row_dimensions[8].height = 22
    sheet.row_dimensions[table_start].height = 23

    # Keep table headings visible while scrolling
    sheet.freeze_panes = "A17"

    # Set print layout
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "portrait"

    # --------------------------------------------------
    # SAVE EXCEL FILE IN MEMORY
    # --------------------------------------------------

    excel_file = BytesIO()

    workbook.save(excel_file)

    excel_file.seek(0)

    # Prepare filename
    filename = (
        f"{class_room.year_level}_"
        f"{class_room.section}_"
        f"Attendance_Report_"
        f"{selected_date.isoformat()}.xlsx"
    )

    filename = filename.replace(" ", "_")

    # Download Excel file
    return send_file(
        excel_file,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# --------------------------------------------------
# STUDENT ATTENDANCE SUMMARY
# --------------------------------------------------

@app.route("/classes/<int:class_id>/student-attendance-summary")
@login_required
def student_attendance_summary(class_id):

    from datetime import date

    class_room = ClassRoom.query.filter_by(
        id=class_id,
        teacher_id=current_user.id
    ).first_or_404()

    # Default date range: current month
    today = date.today()
    default_start = today.replace(day=1)

    start_date_text = request.args.get(
        "start_date",
        default_start.isoformat()
    )

    end_date_text = request.args.get(
        "end_date",
        today.isoformat()
    )

    # Convert dates
    try:

        start_date = date.fromisoformat(
            start_date_text
        )

        end_date = date.fromisoformat(
            end_date_text
        )

    except ValueError:

        flash("Please select valid dates.", "error")

        return redirect(
            url_for(
                "student_attendance_summary",
                class_id=class_id
            )
        )

    if start_date > end_date:

        flash(
            "The start date cannot be after the end date.",
            "error"
        )

        return redirect(
            url_for(
                "student_attendance_summary",
                class_id=class_id
            )
        )

    # Get students
    students = Student.query.filter_by(
        class_id=class_id
    ).all()

    students.sort(
        key=lambda student: (
            student.full_name.split()[-1].lower(),
            student.full_name.lower()
        )
    )

    student_ids = [
        student.id for student in students
    ]

    # Get attendance records within the date range
    if student_ids:

        attendance_records = Attendance.query.filter(
            Attendance.student_id.in_(student_ids),
            Attendance.attendance_date >= start_date,
            Attendance.attendance_date <= end_date
        ).all()

    else:

        attendance_records = []

    # Calculate each student's summary
    summary_rows = []

    for student in students:

        present_count = 0
        absent_count = 0
        late_count = 0

        for record in attendance_records:

            if record.student_id == student.id:

                if record.status == "Present":
                    present_count += 1

                elif record.status == "Absent":
                    absent_count += 1

                elif record.status == "Late":
                    late_count += 1

        recorded_count = (
            present_count + absent_count + late_count
        )

        # Present and Late count as attendance
        if recorded_count > 0:

            attendance_percentage = (
                (present_count + late_count)
                / recorded_count
            ) * 100

        else:

            attendance_percentage = 0

        summary_rows.append({
            "student": student,
            "present": present_count,
            "absent": absent_count,
            "late": late_count,
            "recorded": recorded_count,
            "percentage": attendance_percentage
        })

    return render_template(
        "student_attendance_summary.html",
        class_room=class_room,
        summary_rows=summary_rows,
        start_date=start_date.isoformat(),
        end_date=end_date.isoformat()
    )


# --------------------------------------------------
# DOWNLOAD STUDENT ATTENDANCE SUMMARY AS EXCEL
# --------------------------------------------------

@app.route(
    "/classes/<int:class_id>/student-attendance-summary/export"
)
@login_required
def export_student_attendance_summary(class_id):

    from datetime import date

    # Check class ownership
    class_room = ClassRoom.query.filter_by(
        id=class_id,
        teacher_id=current_user.id
    ).first_or_404()

    # Get selected date range
    start_date_text = request.args.get("start_date", "")
    end_date_text = request.args.get("end_date", "")

    try:

        start_date = date.fromisoformat(start_date_text)
        end_date = date.fromisoformat(end_date_text)

    except ValueError:

        flash("Please select valid dates.", "error")

        return redirect(
            url_for(
                "student_attendance_summary",
                class_id=class_id
            )
        )

    if start_date > end_date:

        flash(
            "The start date cannot be after the end date.",
            "error"
        )

        return redirect(
            url_for(
                "student_attendance_summary",
                class_id=class_id
            )
        )

    # Get students
    students = Student.query.filter_by(
        class_id=class_id
    ).all()

    students.sort(
        key=lambda student: (
            student.full_name.split()[-1].casefold(),
            student.full_name.casefold()
        )
    )

    student_ids = [
        student.id for student in students
    ]

    # Get attendance records within the date range
    if student_ids:

        attendance_records = Attendance.query.filter(
            Attendance.student_id.in_(student_ids),
            Attendance.attendance_date >= start_date,
            Attendance.attendance_date <= end_date
        ).all()

    else:

        attendance_records = []

    # Calculate attendance summary
    summary_rows = []

    for student in students:

        present_count = 0
        absent_count = 0
        late_count = 0

        for record in attendance_records:

            if record.student_id == student.id:

                if record.status == "Present":
                    present_count += 1

                elif record.status == "Absent":
                    absent_count += 1

                elif record.status == "Late":
                    late_count += 1

        recorded_count = (
            present_count + absent_count + late_count
        )

        if recorded_count > 0:

            attendance_percentage = (
                (present_count + late_count)
                / recorded_count
            ) * 100

        else:

            attendance_percentage = 0

        summary_rows.append({
            "name": student.full_name,
            "present": present_count,
            "absent": absent_count,
            "late": late_count,
            "recorded": recorded_count,
            "percentage": attendance_percentage
        })

    # Create Excel workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance Summary"

    # Excel styles
    navy_fill = PatternFill(
        fill_type="solid",
        fgColor="172B4D"
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="21845A"
    )

    light_fill = PatternFill(
        fill_type="solid",
        fgColor="F4F7FB"
    )

    white_font = Font(
        bold=True,
        color="FFFFFF"
    )

    title_font = Font(
        bold=True,
        size=14,
        color="FFFFFF"
    )

    label_font = Font(
        bold=True,
        color="172B4D"
    )

    center_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="DCE4ED"
        )
    )

    # --------------------------------------------------
    # REPORT TITLE
    # --------------------------------------------------

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "STUDENT ATTENDANCE MONITORING SYSTEM"

    sheet["A1"].font = title_font
    sheet["A1"].fill = navy_fill
    sheet["A1"].alignment = center_alignment

    sheet.row_dimensions[1].height = 30

    # --------------------------------------------------
    # CLASS INFORMATION
    # --------------------------------------------------

    sheet["A3"] = "Class"
    sheet["B3"] = (
        f"{class_room.year_level} - {class_room.section}"
    )

    sheet["A4"] = "Subject"
    sheet["B4"] = class_room.subject

    sheet["A5"] = "Academic Year"
    sheet["B5"] = class_room.academic_year

    sheet["A6"] = "Start Date"
    sheet["B6"] = start_date.strftime("%B %d, %Y")

    sheet["A7"] = "End Date"
    sheet["B7"] = end_date.strftime("%B %d, %Y")

    for row in range(3, 8):
        sheet.cell(row=row, column=1).font = label_font

    # --------------------------------------------------
    # SUMMARY TABLE HEADING
    # --------------------------------------------------

    sheet.merge_cells("A9:F9")
    sheet["A9"] = "STUDENT ATTENDANCE SUMMARY"

    sheet["A9"].font = white_font
    sheet["A9"].fill = green_fill
    sheet["A9"].alignment = center_alignment

    # --------------------------------------------------
    # COLUMN HEADINGS
    # --------------------------------------------------

    headers = [
        "No.",
        "Student Name",
        "Present",
        "Absent",
        "Late",
        "Attendance Percentage"
    ]

    header_row = 11

    for column, heading in enumerate(headers, start=1):

        cell = sheet.cell(
            row=header_row,
            column=column,
            value=heading
        )

        cell.font = white_font
        cell.fill = navy_fill
        cell.alignment = center_alignment
        cell.border = thin_border

    # --------------------------------------------------
    # STUDENT RECORDS
    # --------------------------------------------------

    for index, row_data in enumerate(
        summary_rows,
        start=1
    ):

        row_number = header_row + index

        values = [
            index,
            row_data["name"],
            row_data["present"],
            row_data["absent"],
            row_data["late"],
            row_data["percentage"] / 100
        ]

        for column, value in enumerate(values, start=1):

            cell = sheet.cell(
                row=row_number,
                column=column,
                value=value
            )

            cell.border = thin_border

            if index % 2 == 0:
                cell.fill = light_fill

            if column in [1, 3, 4, 5, 6]:
                cell.alignment = center_alignment

            if column == 6:

                cell.font = Font(
                    bold=True,
                    color="21845A"
                )

                cell.number_format = "0.0%"

    # --------------------------------------------------
    # EXCEL FORMATTING
    # --------------------------------------------------

    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 35
    sheet.column_dimensions["C"].width = 15
    sheet.column_dimensions["D"].width = 15
    sheet.column_dimensions["E"].width = 15
    sheet.column_dimensions["F"].width = 25

    sheet.row_dimensions[9].height = 23
    sheet.row_dimensions[11].height = 25

    # Keep headings visible while scrolling
    sheet.freeze_panes = "A12"

    # Print settings
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"

    # --------------------------------------------------
    # SAVE EXCEL FILE
    # --------------------------------------------------

    excel_file = BytesIO()

    workbook.save(excel_file)

    excel_file.seek(0)

    filename = (
        f"{class_room.year_level}_"
        f"{class_room.section}_"
        f"Attendance_Summary_"
        f"{start_date.isoformat()}_"
        f"{end_date.isoformat()}.xlsx"
    )

    filename = filename.replace(" ", "_")

    return send_file(
        excel_file,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )



# --------------------------------------------------
# TEACHER LOGOUT
# --------------------------------------------------

@app.route("/logout", methods=["POST"])
@login_required
def logout():

    logout_user()

    flash("You have logged out.", "success")

    return redirect(url_for("home"))


# --------------------------------------------------
# TEST DATABASE CONNECTION
# --------------------------------------------------

@app.route("/test-db")
def test_db():

    try:

        db.session.execute(
            db.text("SELECT 1")
        )

        return "PostgreSQL connection successful!"

    except Exception:

        db.session.rollback()

        app.logger.exception(
            "Database connection failed"
        )

        return (
            "Database connection failed. Check the terminal.",
            500
        )


# --------------------------------------------------
# RUN APPLICATION
# --------------------------------------------------

if __name__ == "__main__":

    with app.app_context():
        db.create_all()

    app.run(debug=True)
